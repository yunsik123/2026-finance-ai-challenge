"""Build the editable Word specification and its two code-drawn diagrams."""
from pathlib import Path
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'docs/spec-assets'
ASSETS.mkdir(parents=True, exist_ok=True)
INK, TEAL, MUTED = '#173441', '#087E83', '#566B76'
FONT = '/System/Library/Fonts/AppleSDGothicNeo.ttc'


def diagram(name, height, painter):
    canvas = Image.new('RGB', (1600, height), 'white')
    draw = ImageDraw.Draw(canvas)

    def label(x, y, text, size=29, color=INK, anchor='mm'):
        draw.multiline_text((x, y), text, font=ImageFont.truetype(FONT, size),
                            fill=color, anchor=anchor, align='center', spacing=9)

    def box(x, y, w, h, text, fill='#EFF7F7', outline='#C8DEDF', size=29):
        draw.rounded_rectangle((x, y, x+w, y+h), radius=17, fill=fill, outline=outline, width=2)
        label(x+w/2, y+h/2, text, size)

    def arrow(x1, y1, x2, y2, text=None):
        draw.line((x1, y1, x2, y2), fill=TEAL, width=4)
        if x2 > x1:
            draw.polygon([(x2, y2), (x2-13, y2-8), (x2-13, y2+8)], fill=TEAL)
        else:
            draw.polygon([(x2, y2), (x2-8, y2-13), (x2+8, y2-13)], fill=TEAL)
        if text:
            label((x1+x2)/2, (y1+y2)/2-22, text, 22, MUTED)

    painter(label, box, arrow)
    canvas.save(ASSETS / name, dpi=(250, 250))


def flow(label, box, arrow):
    label(110, 93, '소상공인', 33, TEAL)
    for x, text in [(250, '4단계 신청'), (580, '자료 확인·예비평가'), (960, '관리자 검토·승인')]:
        box(x, 43, 290, 100, text, size=28)
    arrow(540, 93, 580, 93)
    arrow(870, 93, 960, 93)
    arrow(1105, 143, 1105, 177)
    box(250, 177, 1000, 64, '승인 후 식당·펀딩 공개', fill='#E4F3F1', size=27)
    arrow(395, 241, 395, 290)
    label(110, 340, '투자자', 33, TEAL)
    for x, text in [(250, '식당·위험정보 확인'), (650, '투자·예약 거래'), (960, '쿠폰 발급·사용·교환')]:
        box(x, 290, 290, 100, text, size=27)
    arrow(540, 340, 650, 340)
    arrow(940, 340, 960, 340)
    label(800, 447, '사장님은 모집 현황과 쿠폰 부담을 관리하고, 투자자는 할인 혜택을 매장에서 이용한다.', 25, MUTED)


def graph(label, box, arrow):
    box(50, 25, 1500, 70, '', fill='#173441', outline='#173441')
    label(800, 60, '질문 예시  “내 신청이 왜 멈췄고, 무엇을 더 준비해야 하나요?”', 30, 'white')
    box(50, 180, 285, 105, '내 신청상태\n검토 대기')
    box(640, 180, 285, 105, '제출자료 요건\n확보·누락 확인')
    box(1240, 180, 310, 105, '자료 발급처\n제출 준비 안내')
    arrow(335, 231, 640, 231, '필요한 자료 확인')
    arrow(925, 231, 1240, 231, '발급 경로 연결')
    label(800, 132, 'Neo4j 관계 검색 · 질문과 관련된 항목에서 최대 2단계 탐색', 26, TEAL)
    arrow(780, 285, 780, 350)
    box(400, 350, 760, 85, '검색 근거 → AI 설명 → 출처·다음 행동 안내', fill='#E4F3F1', size=30)
    label(800, 479, '현재 잔액·주문·심사 현황은 Cloud SQL 원장에서 서버가 직접 집계해 답변', 26, MUTED)


diagram('service-flow.png', 485, flow)
diagram('graph-rag.png', 520, graph)

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(17), Mm(16)
section.left_margin = section.right_margin = Mm(18)
section.header_distance = section.footer_distance = Mm(8)


def font_style(style, size, bold=False, color=INK):
    style.font.name = 'Apple SD Gothic Neo'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), 'Apple SD Gothic Neo')


font_style(doc.styles['Normal'], 10.3)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = Pt(15.2)
normal.space_after = Pt(6)
normal.widow_control = True
font_style(doc.styles['Title'], 26, True)
font_style(doc.styles['Subtitle'], 10.5, False, MUTED)
doc.styles['Subtitle'].font.italic = False
font_style(doc.styles['Heading 1'], 18, True)
font_style(doc.styles['Heading 2'], 12, True, TEAL)
font_style(doc.styles['Caption'], 8.5, False, MUTED)
for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    doc.styles[style_name].element.get_or_add_pPr().append(snap)
for name in ['Heading 1', 'Heading 2']:
    fmt = doc.styles[name].paragraph_format
    fmt.space_before, fmt.space_after = Pt(12), Pt(7)
    fmt.keep_with_next = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(2)
doc.styles['Caption'].paragraph_format.space_after = Pt(8)

head = section.header.paragraphs[0]
head.text = '먹투  /  기능명세서'
head.runs[0].font.size = Pt(8)
head.runs[0].font.color.rgb = RGBColor.from_string('566B76')
foot = section.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = foot.add_run('MEOKTU   ·   ')
run.font.size = Pt(8)
for field in ['PAGE', 'NUMPAGES']:
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), field)
    foot._p.append(fld)
    if field == 'PAGE':
        foot.add_run(' / ')


def inline(paragraph, text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else part.strip('`'))
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.size = Pt(8.5)


def table(lines):
    rows = [[v.strip() for v in line.strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r'[-: ]+', v) for v in r)]
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [32, 79, 63] if '대상·기능' in rows[0] else [35, 75, 64]
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '173441' if ri == 0 else ('F0F6F6' if ri % 2 else 'FAFCFC'))
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge, value_twips in [('top', 74), ('bottom', 74), ('left', 100), ('right', 100)]:
                item = OxmlElement('w:'+edge)
                item.set(qn('w:w'), str(value_twips))
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            tcpr.append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(12.5)
            inline(p, value)
            for r in p.runs:
                r.font.size = Pt(9.2)
                if ri == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                elif ci == 0:
                    r.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.add_run().font.size = Pt(3)


lines = (ROOT / '기능명세서.md').read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        block = []
        while i < len(lines) and lines[i].startswith('|'):
            block.append(lines[i])
            i += 1
        table(block)
        continue
    if line == '<!-- pagebreak -->':
        doc.add_page_break()
    elif line.startswith('!['):
        match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT / match[2]), width=Mm(174))
        cap = doc.add_paragraph(match[1], 'Caption')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    elif i == 2:
        doc.add_paragraph(line, 'Subtitle')
    else:
        p = doc.add_paragraph()
        if line.startswith('- '):
            p.paragraph_format.left_indent = Mm(3)
            p.paragraph_format.first_line_indent = Mm(-3)
            line = '•  ' + line[2:]
        inline(p, line)
    i += 1

doc.core_properties.title = '먹투 기능명세서'
doc.core_properties.subject = 'MVP 기능, 사용자 흐름, AI·GraphDB 처리 및 검증 방법'
doc.core_properties.author = '먹투'
doc.core_properties.keywords = '먹투, MVP, 기능명세서, GraphRAG'
target = ROOT / '기능명세서.docx'
doc.save(target)
loaded = Document(target)
headings = [p.text for p in loaded.paragraphs if p.style.name == 'Heading 1']
assert len(headings) == 5, headings
assert len(loaded.inline_shapes) == 2
assert len(loaded.tables) == 5
print(f'Created {target.name}: 5 main sections, 2 diagrams, 5 tables')
